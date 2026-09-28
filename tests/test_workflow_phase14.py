"""Real-file workflow and resumable native edit checks."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
for path in (ROOT / "runtime" / "src", ROOT / "skill" / "figureloom" / "scripts"):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from figureloom_engine.models import EngineError  # noqa: E402
from figureloom_engine.workflow import (  # noqa: E402
    edit_session,
    normalize_edit,
    parse_edit_phrase,
    preview,
    render_confirmed,
)

FIXTURES = ROOT / "tests" / "fixtures" / "phase14"
CLI = ROOT / "skill" / "figureloom" / "scripts" / "figureloom.py"


def _cli(*args: str, success: bool = True) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(  # noqa: S603 - fixed local Python/CLI, no shell
        [sys.executable, str(CLI), *args], capture_output=True, text=True,
        encoding="utf-8", errors="replace", env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        check=False,
    )
    assert (result.returncode == 0) is success, result.stderr[-1200:]
    return result


def test_preview_profile_confirmation_and_edit_allowlist(tmp_path: Path) -> None:
    source = FIXTURES / "experiment_multiseries.csv"
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    result = preview(source, tmp_path / "run", engine_name="grapher", template_id="trend",
                     engine_home=ROOT / "runtime")
    assert result["source"]["sha256"] == before
    assert result["dataset_profile"]["table"]["row_count"] == 5
    assert {item["name"] for item in result["dataset_profile"]["columns"]} == {
        "X", "Control", "Treatment"
    }
    assert result["recommendation"]["selected_template_id"] == "trend"
    assert result["understanding"]["confirmation_gate"]["can_confirm_now"]
    with pytest.raises(EngineError, match="Confirm") as error:
        render_confirmed(tmp_path / "run" / "workflow-preview.json", claim="comparison")
    assert error.value.code == "confirmation_required"
    assert parse_edit_phrase("把 Treatment 改成虚线") == {
        "operation": "set_line_style", "series": "Treatment", "value": "dashed"
    }
    assert parse_edit_phrase("把 Y 轴标题改成 Current") == {
        "operation": "set_axis_title", "axis": "y", "value": "Current"
    }
    with pytest.raises(EngineError) as error:
        normalize_edit("raw_script", value="Origin C")
    assert error.value.code == "edit_unsupported"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    automatic = preview(source, tmp_path / "automatic", engine_name="grapher",
                        engine_home=ROOT / "runtime")
    assert automatic["recommendation"]["selected_template_id"] == "trend"


def test_xlsx_selected_sheet_and_missing_artifact(tmp_path: Path) -> None:
    from openpyxl import Workbook

    source = tmp_path / "experiment.xlsx"
    book = Workbook()
    book.active.title = "Notes"
    book.active.append(["Ignore this sheet"])
    data = book.create_sheet("Data")
    data.append(["X", "Control", "Control_SD", "Treatment", "Treatment_SD"])
    for number in range(1, 6):
        data.append([number, number * 2, 0.2, number * 2.3, 0.3])
    book.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    result = preview(source, tmp_path / "run", engine_name="origin", template_id="line_error",
                     sheet="Data", engine_home=ROOT / "runtime")
    assert result["source"]["selected_sheet"] == "Data"
    assert result["source"]["sheets"] == ["Notes", "Data"]
    assert result["effective_source"]["path"].endswith("selected_sheet.csv")
    assert result["dataset_profile"]["table"]["row_count"] == 5
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    fake = tmp_path / "session.json"
    fake.write_text(json.dumps({
        "status": "verified", "artifacts": {"editable": str(tmp_path / "missing.opju")}
    }))
    with pytest.raises(EngineError) as error:
        edit_session(fake, {"operation": "set_axis_title", "axis": "y", "value": "Current"})
    assert error.value.code == "artifact_not_found"


def test_source_change_and_sheet_errors_are_structured(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    source.write_text("X,Control,Treatment\n1,2,3\n2,4,6\n3,6,9\n", encoding="utf-8")
    with pytest.raises(EngineError) as error:
        preview(source, tmp_path / "bad-sheet", engine_name="grapher", sheet="Data")
    assert error.value.code == "sheet_unsupported"
    preview(source, tmp_path / "run", engine_name="grapher", template_id="trend",
            engine_home=ROOT / "runtime")
    source.write_text(source.read_text(encoding="utf-8") + "4,8,12\n", encoding="utf-8")
    with pytest.raises(EngineError) as error:
        render_confirmed(tmp_path / "run" / "workflow-preview.json", claim="changed", confirmed=True)
    assert error.value.code == "source_changed_since_preview"
    assert not (tmp_path / "run" / "session.json").exists()


def test_missing_or_invalid_workflow_files_return_error_codes(tmp_path: Path) -> None:
    missing = tmp_path / "missing.json"
    with pytest.raises(EngineError) as error:
        render_confirmed(missing, claim="comparison", confirmed=True)
    assert error.value.code == "preview_not_found"
    with pytest.raises(EngineError) as error:
        edit_session(missing, {"operation": "set_axis_title", "axis": "y", "value": "Current"})
    assert error.value.code == "session_not_found"
    missing.write_text("[]", encoding="utf-8")
    with pytest.raises(EngineError) as error:
        render_confirmed(missing, claim="comparison", confirmed=True)
    assert error.value.code == "preview_invalid"


def test_grapher_explicit_weight_rejected_before_successful_session(tmp_path: Path) -> None:
    source = FIXTURES / "experiment_weighted.csv"
    mapping = {"assignments": {"X": "x", "Y": "series", "W": "ignored"}}
    fit = {"model": "linear", "x_column": "X", "y_column": "Y", "weight_mode": "column",
           "weight_column": "W", "weight_interpretation": "direct_weight"}
    preview(source, tmp_path / "run", engine_name="grapher", template_id="scatter",
            mapping=mapping, fit_spec=fit, engine_home=ROOT / "runtime")
    with pytest.raises(EngineError) as error:
        render_confirmed(tmp_path / "run" / "workflow-preview.json", claim="weighted", confirmed=True)
    assert error.value.code == "unsupported_fit_weighting"
    failed = json.loads((tmp_path / "run" / "session.json").read_text(encoding="utf-8"))
    assert failed["status"] == "failed"
    assert failed["error"]["code"] == "unsupported_fit_weighting"
    assert not list((tmp_path / "run").rglob("*.grf"))


def test_minimal_fit_input_uses_canonical_existing_schema(tmp_path: Path) -> None:
    result = preview(
        FIXTURES / "experiment_fit.csv", tmp_path / "run", engine_name="grapher",
        template_id="scatter", fit_spec={"model": "linear", "x_column": "X", "y_column": "Y"},
        engine_home=ROOT / "runtime",
    )
    assert result["fit_spec"]["weight_mode"] == "none"
    assert result["fit_spec"]["result_source"] == "backend_native"


@pytest.mark.grapher
@pytest.mark.skipif(os.name != "nt", reason="Grapher COM requires Windows")
def test_csv_grapher_render_resume_and_native_series_edit(tmp_path: Path) -> None:
    from grapher_sciplot.smoke import SmokeFailure, discover, grapher_pids

    try:
        discover()
    except SmokeFailure as exc:
        pytest.skip(str(exc))
    before = grapher_pids()
    run = tmp_path / "run"
    source = FIXTURES / "experiment_multiseries.csv"
    _cli("workflow-preview", str(source), "--engine", "grapher", "--template-id", "trend",
         "--output-dir", str(run), "--engine-home", str(ROOT / "runtime"))
    _cli("workflow-render", str(run / "workflow-preview.json"), "--claim", "Compare treatments",
         "--confirm")
    session = json.loads((run / "session.json").read_text(encoding="utf-8"))
    assert session["status"] == "verified"
    assert Path(session["artifacts"]["editable"]).is_file()
    assert all(Path(value).is_file() for value in session["artifacts"]["exports"].values())
    edit_output = _cli("edit", str(run / "session.json"), "把 Treatment 改成虚线")
    session = json.loads((run / "session.json").read_text(encoding="utf-8"))
    assert session["edits"][-1]["verification"]["actual"] == ".1 in. Dash"
    native = json.loads(edit_output.stdout)["native"]["readback"]
    assert {plot["name"]: plot["line_style"] for plot in native["plots"]} == {
        "Control": "Solid", "Treatment": ".1 in. Dash"
    }
    assert grapher_pids() == before


@pytest.mark.origin
@pytest.mark.skipif(os.name != "nt", reason="Origin automation requires Windows")
def test_xlsx_origin_render_resume_and_native_axis_edit(tmp_path: Path) -> None:
    import figureloom_core as core
    from openpyxl import Workbook

    if not core.discover_origin_application()["launch_registration_detected"]:
        pytest.skip("Origin is not registered")
    source = tmp_path / "experiment.xlsx"
    book = Workbook()
    book.active.title = "Notes"
    book.active.append(["Not data"])
    data = book.create_sheet("Data")
    data.append(["X", "Control", "Control_SD", "Treatment", "Treatment_SD"])
    for number in range(1, 6):
        data.append([number, number * 2, 0.2, number * 2.3, 0.3])
    book.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    run = tmp_path / "run"
    _cli("workflow-preview", str(source), "--engine", "origin", "--template-id", "line_error",
         "--sheet", "Data", "--output-dir", str(run), "--engine-home", str(ROOT / "runtime"))
    _cli("workflow-render", str(run / "workflow-preview.json"), "--claim", "Compare error bars",
         "--confirm")
    session = json.loads((run / "session.json").read_text(encoding="utf-8"))
    assert all(Path(value).is_file() for value in session["artifacts"]["exports"].values())
    _cli("edit", str(run / "session.json"), "把 Y 轴标题改成 Current (mA)")
    session = json.loads((run / "session.json").read_text(encoding="utf-8"))
    assert session["edits"][-1]["verification"]["actual"] == "Current (mA)"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before


@pytest.mark.origin
@pytest.mark.skipif(os.name != "nt", reason="Both native engines require Windows")
def test_same_confirmed_plan_runs_on_both_engines(tmp_path: Path) -> None:
    import figureloom_core as core
    from grapher_sciplot.smoke import SmokeFailure, discover

    if not core.discover_origin_application()["launch_registration_detected"]:
        pytest.skip("Origin is not registered")
    try:
        discover()
    except SmokeFailure as exc:
        pytest.skip(str(exc))
    run = tmp_path / "run"
    source = FIXTURES / "experiment_multiseries.csv"
    _cli("workflow-preview", str(source), "--engine", "grapher", "--template-id", "trend",
         "--output-dir", str(run), "--engine-home", str(ROOT / "runtime"))
    _cli("workflow-render", str(run / "workflow-preview.json"), "--claim", "Same plan",
         "--confirm")
    _cli("render", str(run / "render-plan.json"), "--engine", "origin",
         "--engine-home", str(ROOT / "runtime"), "--output-dir", str(run / "origin"),
         "--close-origin")
    plan = json.loads((run / "render-plan.json").read_text(encoding="utf-8"))
    grapher = json.loads((run / "grapher" / "grapher_verify_report.json").read_text(encoding="utf-8"))
    origin = json.loads((run / "origin" / "origin_verify_report.json").read_text(encoding="utf-8"))
    expected = plan["render_spec"]["data"]["y"]
    assert [plot["name"] for plot in grapher["readback"]["plots"]] == expected
    assert [item["source_column"] for item in origin["origin_series_columns"]] == expected
    assert all((run / "origin" / f"result.{suffix}").is_file()
               for suffix in ("opju", "png", "pdf", "tif"))
